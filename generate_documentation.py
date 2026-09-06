from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "ASG_Airlines_Data_Pipeline_Documentation.docx"
ASSET = ROOT / "documentation_assets"
ASSET.mkdir(exist_ok=True)

NAVY = "16324F"
BLUE = "1479E8"
TEAL = "1B9AAA"
AMBER = "F2A900"
LIGHT = "F3F7FB"
GREY = "5B6770"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF', 9); shade(t.rows[0].cells[i], NAVY)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, False, None, 8.5); shade(cells[i], 'FFFFFF' if ridx % 2 else LIGHT)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def add_box(draw, xy, text, fill, outline=NAVY, font=None):
    fill = '#' + fill if isinstance(fill, str) and not fill.startswith('#') else fill
    outline = '#' + outline if isinstance(outline, str) and not outline.startswith('#') else outline
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=3)
    x1,y1,x2,y2=xy
    lines=[]
    for line in text.split('\n'):
        while len(line)>22:
            cut=line.rfind(' ',0,22)
            if cut<1: cut=22
            lines.append(line[:cut]); line=line[cut:].strip()
        lines.append(line)
    y=y1+(y2-y1-len(lines)*18)//2
    for line in lines:
        bb=draw.textbbox((0,0),line,font=font); x=(x1+x2-(bb[2]-bb[0]))//2
        draw.text((x,y),line,fill='#'+NAVY,font=font); y+=18

def arrow(draw, a, b, fill=BLUE):
    fill = '#' + fill if isinstance(fill, str) and not fill.startswith('#') else fill
    draw.line([a,b], fill=fill, width=5)
    import math
    ang=math.atan2(b[1]-a[1], b[0]-a[0]); L=13
    pts=[b,(b[0]-L*math.cos(ang-0.45),b[1]-L*math.sin(ang-0.45)),(b[0]-L*math.cos(ang+0.45),b[1]-L*math.sin(ang+0.45))]
    draw.polygon(pts, fill=fill)

def make_architecture():
    im=Image.new('RGB',(1200,430),'white'); d=ImageDraw.Draw(im); f=ImageFont.load_default()
    add_box(d,(30,155,210,275),'Excel source\n4 logical entities','DDEBFA',font=f)
    add_box(d,(275,155,455,275),'ADF / batch\ningestion','E5F2F4',font=f)
    add_box(d,(520,80,735,175),'Bronze / Raw\nimmutable + audit','FFF2CC',font=f)
    add_box(d,(520,255,735,350),'Databricks / Python\nprofile / validate / transform','D9EAF7',font=f)
    add_box(d,(800,80,1015,175),'Silver / Curated\nPII-safe + quarantine','E2F0D9',font=f)
    add_box(d,(800,255,1015,350),'Gold / serving\nfacts / dimensions / KPIs','D9EAF7',font=f)
    add_box(d,(1080,175,1180,275),'Power BI\nreport','DDEBFA',font=f)
    arrow(d,(210,215),(275,215)); arrow(d,(455,215),(520,128)); arrow(d,(455,215),(520,300)); arrow(d,(735,128),(800,128)); arrow(d,(735,300),(800,300)); arrow(d,(1015,300),(1080,225)); arrow(d,(1015,128),(1080,225))
    im.save(ASSET/'architecture.png')

def make_dataflow():
    im=Image.new('RGB',(1200,360),'white'); d=ImageDraw.Draw(im); f=ImageFont.load_default()
    labels=[('1. Ingest','Workbook\nschema checks'),('2. Profile','counts, nulls,\nIDs, timestamps'),('3. Silver','clean text, parse\ntimes, hash PII'),('4. Validate','quality flags +\nquarantine'),('5. Gold','facts, dims,\naggregates'),('6. Consume','Power BI\ninteractive pages')]
    xs=[20,215,410,605,800,995]
    for x,(title,body) in zip(xs,labels):
        add_box(d,(x,125,x+165,235),title+'\n'+body,'EAF3FA' if x<800 else 'E2F0D9',font=f)
    for x in xs[:-1]: arrow(d,(x+165,180),(x+190,180))
    d.text((50,40),'Rejected rows -> reason-coded quarantine (parallel audit path)',fill='#'+AMBER,font=f)
    d.line((600,55,600,125),fill='#'+AMBER,width=4); arrow(d,(600,55),(600,125),AMBER)
    im.save(ASSET/'dataflow.png')

def make_model():
    im=Image.new('RGB',(1200,430),'white'); d=ImageDraw.Draw(im); f=ImageFont.load_default()
    add_box(d,(435,160,765,280),'fact_flight_operations\nflight_id / route / airline\ndeparture/arrival / duration','DDEBFA',font=f)
    add_box(d,(60,40,300,135),'dim_airline\nairline','E2F0D9',font=f)
    add_box(d,(60,295,300,390),'dim_route\nsource / destination / route','E2F0D9',font=f)
    add_box(d,(900,40,1140,135),'dim_date\ndeparture_date','E2F0D9',font=f)
    add_box(d,(900,295,1140,390),'fact_bookings\nbooking_id / status / flight_id','FFF2CC',font=f)
    add_box(d,(435,330,765,420),'fact_payments\npayment_id / booking_id / amount','FFF2CC',font=f)
    for a,b in [((300,90),(435,190)),((300,340),(435,245)),((900,90),(765,190)),((900,340),(765,250)),((765,345),(900,345))]: arrow(d,a,b,TEAL)
    im.save(ASSET/'data_model.png')

def heading(doc, text, level=1):
    p=doc.add_heading(text, level=level); return p

def bullet(doc, text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def code(doc, text):
    p=doc.add_paragraph(); p.style='Code'; r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8); return p

def add_picture(doc, path, width=6.5, caption=None):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
    if caption:
        q=doc.add_paragraph(caption); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.runs[0].italic=True; q.runs[0].font.size=Pt(8); q.runs[0].font.color.rgb=RGBColor.from_string(GREY)

def main():
    make_architecture(); make_dataflow(); make_model()
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor.from_string('263238')
    for s in ['Title','Heading 1','Heading 2','Heading 3']:
        styles[s].font.name='Aptos Display'; styles[s].font.color.rgb=RGBColor.from_string(NAVY)
    styles['Heading 1'].font.size=Pt(18); styles['Heading 2'].font.size=Pt(13); styles['Heading 3'].font.size=Pt(10.5)
    if 'Code' not in [s.name for s in styles]:
        st=styles.add_style('Code',WD_STYLE_TYPE.PARAGRAPH); st.font.name='Consolas'; st.font.size=Pt(8); st.paragraph_format.left_indent=Inches(.2)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('ASG AIRLINES'); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(BLUE)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('End-to-End Flight Operations Data Pipeline'); r.bold=True; r.font.size=Pt(25); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Technical documentation, data model, quality strategy and Power BI reporting handoff'); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(GREY)
    doc.add_paragraph(); add_picture(doc,ROOT/'Images'/'Executive_Overview.png',3.0,'Power BI Executive Overview — implemented report page')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(f'Prepared {datetime.now():%d %B %Y}  |  Repository: airlines-case-study'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GREY)
    doc.add_page_break()
    heading(doc,'1. Executive summary')
    doc.add_paragraph('This submission implements a reproducible local batch pipeline that ingests the supplied Excel workbook, profiles source quality, creates cleaned and PII-safe Silver outputs, quarantines rejected records, builds Gold facts/dimensions/KPI aggregates, and provides a Power BI report design with five operational pages. The implementation uses Python and pandas locally; the target production architecture maps the same logic to Azure Data Factory, ADLS Gen2, Databricks/Delta, Azure Monitor and Power BI.')
    table(doc,['Outcome','Implemented result'],[
        ('Functional pipeline','profile.py → transform.py → gold.py; CSV and Parquet outputs'),
        ('Duration correctness','timestamp parsing, next-day rollover, recalculated minutes, anomaly flags'),
        ('Data quality','explicit rules, reason-coded quarantine, source and Silver quality CSVs'),
        ('Privacy','raw passenger attributes removed; passenger_id represented only as SHA-256 hash'),
        ('Reporting','PBIX plus five page captures and DAX measure catalogue'),
    ])
    heading(doc,'2. Repository implementation and run instructions')
    code(doc,'uv sync\nuv run pytest\nuv run python src/profile.py\nuv run python src/transform.py\nuv run python src/gold.py')
    doc.add_paragraph('The source workbook is data/raw/UseCase - Airlines.xlsx. The scripts are intentionally rerunnable: each run reads the source, recreates curated and Gold outputs, and writes audit-friendly CSV/Parquet files. Tests cover overnight duration adjustment, PII removal from passenger and booking outputs, and Gold KPI reconciliation.')
    table(doc,['Component','Responsibility','Output'],[
        ('src/profile.py','Schema and source profiling before cleaning','data/quality/source_profile.csv'),
        ('src/transform.py','Normalize, type, validate, hash PII, quarantine','data/curated/*; data/quarantine/records.csv; Silver quality'),
        ('src/gold.py','Filter valid relationships and aggregate','data/gold/* facts, dimensions and KPI tables'),
        ('tests/','Regression checks for core business rules','pytest pass/fail'),
        ('powerbi/','Model guidance and DAX measures','ASG_Airlines.pbix and measures.dax'),
    ])
    heading(doc,'3. Architecture diagram')
    add_picture(doc,ASSET/'architecture.png',6.7,'Figure 1. Azure-first target architecture with equivalent local implementation')
    doc.add_paragraph('Bronze is the immutable evidence layer and must remain restricted. Silver is the cleaned and PII-safe analytical layer. Gold is the serving layer designed for Power BI. In production, ADF owns orchestration and parameterization; ADLS Gen2 stores the medallion layers; Databricks/Delta provides scalable transformations; Entra ID, Key Vault, Unity Catalog and Azure Monitor provide security and operations.')
    heading(doc,'4. Data flow diagram')
    add_picture(doc,ASSET/'dataflow.png',6.7,'Figure 2. End-to-end data flow and rejection path')
    doc.add_paragraph('The rejection path is deliberately parallel to the successful path. Invalid records are not silently discarded: they are retained with dataset and reason fields so data owners can remediate source issues and reconcile counts.')
    heading(doc,'5. Source dataset documentation')
    table(doc,['Sheet / entity','Rows × columns','Primary purpose','Sensitive fields / notes'],[
        ('flights','1,020 × 7','Flight schedule and operational timestamps','Malformed IDs, missing airline values, overnight times'),
        ('bookings','1,000 × 9','Booking-to-flight and passenger relationship','passport and emergency-contact PII removed from safe output'),
        ('payments','1,000 × 4','Payment transactions linked to bookings','missing amounts remain null'),
        ('passengers','1,039 × 9','Passenger attributes','name, email, phone, Aadhaar and DOB are restricted raw PII'),
    ])
    doc.add_paragraph('Original columns are preserved in the raw workbook. Serving tables expose only approved analytical attributes. The Gold model is intentionally operational and aggregate-friendly rather than a customer-360 model.')
    heading(doc,'6. Data model')
    add_picture(doc,ASSET/'data_model.png',6.7,'Figure 3. Gold star/snowflake model and relationship directions')
    doc.add_paragraph('Gold grains: one row per valid flight operation, valid booking, payment row with a valid booking reference, route, airline, and departure date. Relationships should be single-direction and one-to-many. The model avoids many-to-many joins and keeps route and airline dimensions deduplicated for filter integrity.')
    table(doc,['Table','Grain','Important fields'],[
        ('fact_flight_operations','one valid flight','flight_id, airline, source, destination, route, departure_time, arrival_time_adjusted, duration_minutes, overnight_flag'),
        ('fact_bookings','one valid booking','booking_id, passenger_key_hash, flight_id, booking_date, status, seat_number'),
        ('fact_payments','one payment row with valid booking','payment_id, booking_id, amount, payment_method'),
        ('dim_airline','one airline','airline'),('dim_route','one source/destination route','source, destination, route'),('dim_date','one departure date','departure_date'),
        ('agg_*_kpis','one route/airline/executive metric grain','pre-aggregated counts, averages, rates and totals'),
    ])
    heading(doc,'7. Preprocessing and transformation logic')
    heading(doc,'7.1 Common cleaning',2)
    for x in ['Column names are trimmed and lower-cased; text whitespace is removed; blank strings become null.','Dates and timestamps are parsed with coercion so parse failures become explicit quality issues rather than crashing downstream aggregations.','Numeric amounts and source duration values are coerced to numeric; source duration is retained as duration_raw for audit comparison but not used as the KPI value.','Exact duplicates and missing-key conditions are profiled and reported.','Known categories are normalized, including uppercase booking statuses and a visible Unknown airline bucket.']:
        bullet(doc,x)
    heading(doc,'7.2 Flight IDs and validation',2)
    doc.add_paragraph('A valid flight identifier must match the regular expression ^[A-Z]{2}\\d{3}$. Malformed identifiers are upper-cased for consistency but never guessed or rewritten. The raw value is preserved in flight_id_raw, flight_id_valid_flag is false, and the row is sent to quarantine.')
    heading(doc,'7.3 Duration calculation and overnight handling',2)
    code(doc,'arrival_time_adjusted = arrival_time\nif arrival_time < departure_time:\n    arrival_time_adjusted += 1 day\nduration_minutes = (arrival_time_adjusted - departure_time).total_seconds() / 60\novernight_flag = date(arrival_time) > date(departure_time) OR rollover')
    doc.add_paragraph('This rule handles a flight departing late on one calendar day and arriving after midnight on the next. The recalculated duration, not the source duration field, drives Gold KPIs. Null, non-positive or over-24-hour values are flagged as duration anomalies. Example regression case: departure 23:30 and arrival 01:15 becomes 105 minutes after adding one day to arrival.')
    heading(doc,'7.4 Bookings and payments',2)
    for x in ['Allowed statuses are CONFIRMED, CANCELLED and PENDING. Cancelled rows remain available for cancellation-rate reporting; invalid statuses or missing booking keys are quarantined.','Booking flight references are constrained to valid Gold flight IDs. Payment booking references are constrained to valid booking IDs.','Missing payment amounts remain null and are excluded from payment totals; they are never converted to zero.','The passenger relationship uses passenger_key_hash, a one-way SHA-256 hash, so Power BI can relate operational records without exposing passenger identifiers.']:
        bullet(doc,x)
    heading(doc,'8. Data quality, governance and assumptions')
    table(doc,['Rule family','Decision','Operational response'],[
        ('Schema','Required sheets and columns must exist','Fail fast before transformation'),('Identifiers','Validate, do not infer malformed IDs','Flag + quarantine + audit'),('Timestamps','Parse; handle next-day rollover','Recalculate duration and flag anomalies'),('Foreign keys','Bookings/payments must resolve','Exclude from business KPIs; retain rejection evidence'),('PII','Raw passenger/contact data is restricted','Hash approved key; remove raw PII from Silver-safe/Gold/Power BI'),('Missing amounts','Null is different from zero','Exclude from sums and display quality issue'),
    ])
    doc.add_paragraph('Assumptions: the first arrival timestamp earlier than departure represents next-day arrival; no time-zone conversion is performed because the supplied workbook does not provide time zones; a malformed identifier cannot be corrected without an approved reference mapping; KPI denominators are documented and invalid records remain available for audit.')
    heading(doc,'9. Current quality and KPI evidence')
    table(doc,['Metric','Observed result','Interpretation'],[
        ('Valid Gold flights','747','Rows passing flight ID and duration validation'),('Average duration','164.1 minutes','Recalculated adjusted timestamps'),('Median duration','164 minutes','Gold flight operations'),('Overnight flights','93 valid Gold / 125 flagged in source-quality view','Different denominators are intentional'),('Valid bookings','679','Approved status + valid flight reference'),('Cancellation rate','33.7%','Cancelled valid bookings ÷ all valid bookings'),('Payment total','5,419,385.34','Non-null payment amounts with valid booking references'),('Quarantined records','348','Reason-coded audit output'),
    ])
    doc.add_paragraph('The dashboard quality page intentionally exposes both source-quality findings and Gold-valid KPIs. This avoids conflating “found in source” with “eligible for business reporting.”')
    heading(doc,'10. Power BI report and reporting design')
    doc.add_paragraph('The implemented PBIX is ASG_Airlines.pbix. It is designed to load Gold CSVs plus the two quality CSVs; raw and curated passenger data must not be loaded. The report uses slicers for departure date, airline, route, source, destination and overnight flag, with consistent navy/blue/amber visual language.')
    table(doc,['Page','Decision-support question','Key visuals'],[
        ('Executive Overview','What is the current operating picture?','KPI cards, airline distribution, top routes, date trend, slicers'),('Duration Analysis','Where are duration and overnight patterns?','Distribution, average by airline/route, same-day vs overnight, flight detail'),('Route Performance','Which routes carry traffic and bookings?','Rankings, source/destination matrix, booking volume, route duration'),('Airline Trends','How do airlines compare?','Flights, duration, booking status, cancellation rate, payment total'),('Data Quality & Anomalies','Can the numbers be trusted?','Missing values, malformed IDs, quarantine, anomalies, quality failures'),
    ])
    for pth, cap in [('Executive_Overview.png','Executive overview'),('Duration_Analysis.png','Duration analysis'),('Route Performance.png','Route performance'),('Airline_Trends.png','Airline trends'),('Data Quality & Anomalies.png','Data quality and anomalies')]:
        if (ROOT/'Images'/pth).exists(): add_picture(doc,ROOT/'Images'/pth,4.3,cap+' — Power BI page capture')
    heading(doc,'11. KPI catalogue and DAX patterns')
    code(doc,'Total Flights = COUNTROWS(fact_flight_operations)\nAverage Flight Duration (min) = AVERAGE(fact_flight_operations[duration_minutes])\nOvernight Flight Rate = DIVIDE([Overnight Flights], [Total Flights])\nCancellation Rate = DIVIDE([Cancelled Bookings], [Valid Bookings])\nPayment Total = SUM(fact_payments[amount])')
    doc.add_paragraph('Measures use explicit denominators and preserve null payment semantics. Format durations as whole minutes, rates as percentages, and payments in the selected currency. Use aggregates for fast overview visuals and facts for drill-through/detail analysis.')
    heading(doc,'12. Scalability, performance and production hardening')
    for x in ['Replace the local workbook source with parameterized ADF ingestion into ADLS Gen2 and use run IDs, source hashes and immutable raw copies.','Use Databricks/Spark and Delta tables for larger files, partitioning by ingestion date or departure date, and incremental processing using a watermark.','Add data contracts, schema evolution controls, idempotent writes, and automated quality gates that stop Gold publication when critical rules fail.','Use managed identities, Key Vault, Entra groups, ADLS ACLs, Unity Catalog and Power BI row/object permissions as appropriate.','Add Azure Monitor diagnostics, refresh alerts, dead-letter/quarantine retention policy and deployment through GitHub Actions or Azure DevOps.']:
        bullet(doc,x)
    heading(doc,'13. Validation and handoff checklist')
    for x in ['Run pytest and confirm overnight, PII and KPI reconciliation tests pass.','Run profile → transform → gold in order and inspect quality/quarantine outputs.','Refresh Power BI after any schema change; verify relationships, slicers, drill-through and KPI reconciliation against agg_executive_kpis.csv.','Confirm no raw passenger columns appear in the model, visual fields, filters, tooltips or drill-through pages.','Publish only Gold and approved quality outputs; retain raw and quarantine layers in restricted storage.']:
        bullet(doc,x)
    heading(doc,'Appendix A — Key implementation excerpt')
    code(doc,'# src/transform.py\nframe["flight_id_valid_flag"] = frame["flight_id"].fillna("").str.fullmatch(r"[A-Z]{2}\\d{3}")\nneeds_rollover = valid_times & (frame["arrival_time"] < frame["departure_time"])\nframe.loc[needs_rollover, "arrival_time_adjusted"] += pd.Timedelta(days=1)\nframe["duration_minutes"] = (frame["arrival_time_adjusted"] - frame["departure_time"]).dt.total_seconds() / 60\n\n# src/gold.py\nvalid_flights = flights[_flag(flights, "flight_id_valid_flag") & ~_flag(flights, "duration_anomaly_flag")]\nvalid_bookings = bookings[bookings["status"].isin(["CONFIRMED", "CANCELLED", "PENDING"])]')
    doc.add_paragraph('Full source code is retained in src/profile.py, src/transform.py, src/gold.py, with automated checks in tests/. This document is the design and operational handoff layer for the implementation.')
    doc.save(OUT); print(OUT)

if __name__ == '__main__': main()
